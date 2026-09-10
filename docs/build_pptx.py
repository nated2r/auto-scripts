"""Generate 自動產稿機器人 product deck as .pptx (16:9)."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

GREEN = RGBColor(0x06, 0xC7, 0x55)
GREEN_D = RGBColor(0x05, 0xA8, 0x48)
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xE8, 0xF8, 0xEE)
CARD = RGBColor(0xF7, 0xFB, 0xF8)
INK = RGBColor(0x14, 0x20, 0x33)

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size=18, bold=False, color=INK, name="Microsoft JhengHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    if line is not None:
        shape.line.color.rgb = line
    try:
        shape.adjustments[0] = 0.15
    except Exception:
        pass
    return shape


def footer(slide, text, page, total=8):
    bar = add_rect(slide, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.65), NAVY)
    tf = bar.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"  {text}"
    set_run(run, size=14, bold=True, color=WHITE)
    add_textbox(
        slide,
        Inches(11.2),
        Inches(6.65),
        Inches(1.5),
        Inches(0.4),
        f"{page:02d} / {total:02d}",
        size=12,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.RIGHT,
    )


def card(slide, left, top, width, height, title, body, fill=WHITE, title_color=NAVY, body_color=MUTED):
    add_rect(slide, left, top, width, height, fill)
    add_textbox(slide, left + Inches(0.22), top + Inches(0.2), width - Inches(0.4), Inches(0.4), title, size=18, bold=True, color=title_color)
    box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.65), width - Inches(0.4), height - Inches(0.85))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=13, color=body_color)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 01
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), W, H, WHITE)
    add_textbox(s, Inches(0.7), Inches(1.3), Inches(11), Inches(0.4), "SHORT SCRIPT OS", size=16, bold=True, color=GREEN_D)
    add_textbox(s, Inches(0.7), Inches(1.8), Inches(11), Inches(2.2), "貼稿進去\n爆款口播自動出來", size=44, bold=True, color=NAVY)
    add_textbox(
        s,
        Inches(0.7),
        Inches(4.2),
        Inches(10),
        Inches(1.2),
        "自動產稿機器人｜新聞／參考口播／影片轉錄 → IG・FB 約 30 秒、高留言導向六欄稿。\n台灣繁中、結構保底、規則可自定；可加值串接轉錄機器人。",
        size=16,
        color=MUTED,
    )
    footer(s, "你不是在「請 AI 亂寫」——是在裝一台短影音產稿引擎", 1)

    # 02
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), W, H, WHITE)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(10), Inches(0.35), "為什麼需要", size=14, bold=True, color=GREEN_D)
    add_textbox(s, Inches(0.7), Inches(0.85), Inches(11), Inches(1.2), "每天要產稿\n卡的不是靈感，是結構", size=34, bold=True, color=NAVY)
    card(s, Inches(0.7), Inches(2.4), Inches(3.7), Inches(3.3), "1｜寫得出來，但不爆", "有內容、沒鉤子、沒懸念、\nCTA 模糊，留言起不來。", SOFT)
    card(s, Inches(4.7), Inches(2.4), Inches(3.7), Inches(3.3), "2｜每支格式都不一樣", "封面、標題、口播、內文、\n標籤、懶人包東拼西湊，\n貼 Sheet 更痛。")
    card(s, Inches(8.7), Inches(2.4), Inches(3.7), Inches(3.3), "3｜下一篇就感覺哪裡都不對", "沒有帳號規則沉澱，\n產出來不像你、也不穩定。", SOFT)
    footer(s, "結果句：需要的是可複製的產稿系統，不是再一個聊天框。", 2)

    # 03
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), W, H, WHITE)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(10), Inches(0.35), "產品是什麼", size=14, bold=True, color=GREEN_D)
    add_textbox(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1.1), "把「參考稿／題材／影片」變成可上架口播包", size=28, bold=True, color=NAVY)
    card(
        s,
        Inches(0.7),
        Inches(2.2),
        Inches(5.8),
        Inches(3.5),
        "輸入 → 輸出",
        "輸入：新聞／舊口播（改爆款）或一句題材（展開）\n也可：影片網址 → 轉錄文字 → 貼回產稿\n\n輸出六欄：\n封面文案｜影片標題｜口播稿\n內文｜hashtag｜懶人包整理",
        SOFT,
    )
    card(
        s,
        Inches(6.8),
        Inches(2.2),
        Inches(5.6),
        Inches(3.5),
        "工作台規格",
        "約 30–45 秒｜口播 150 字基準\n一次一個角度\n一鍵複製可贴 Google Sheet\n轉錄機器人餵文字進來",
        NAVY,
        title_color=WHITE,
        body_color=RGBColor(0xC9, 0xE6, 0xD4),
    )
    footer(s, "結果句：從「想到寫」變成「貼上就產、對齊就發」。", 3)

    # 04
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), W, H, WHITE)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(10), Inches(0.35), "兩種產稿入口", size=14, bold=True, color=GREEN_D)
    add_textbox(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1.1), "有稿就改爆款　有題就展開角度", size=32, bold=True, color=NAVY)
    card(
        s,
        Inches(0.7),
        Inches(2.2),
        Inches(5.5),
        Inches(3.4),
        "貼稿改爆款",
        "貼新聞／參考口播／舊稿／轉錄文字\n• 萃取題材後整篇重寫\n• 不是逐句微調\n• 適合：已有素材要升級",
        SOFT,
    )
    card(
        s,
        Inches(7.0),
        Inches(2.2),
        Inches(5.5),
        Inches(3.4),
        "題材展開",
        "只丟一句題材名\n• 依角度從零產一支\n• 反直覺／測驗／解讀…\n• 適合：只有題目要開稿",
    )
    footer(s, "結果句：材料不同，引擎相同——都落到同一套六欄規格。", 4)

    # 05
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), W, H, WHITE)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(10), Inches(0.35), "輸出規格", size=14, bold=True, color=GREEN_D)
    add_textbox(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1.0), "一格一格分開，也能單獨複製", size=32, bold=True, color=NAVY)
    card(s, Inches(0.7), Inches(2.2), Inches(3.8), Inches(3.3), "拍攝／剪輯用", "封面文案、影片標題、口播稿\n進棚就能念、能上字幕", NAVY, title_color=WHITE, body_color=RGBColor(0xC9, 0xE6, 0xD4))
    card(s, Inches(4.75), Inches(2.2), Inches(3.8), Inches(3.3), "發文用", "內文四層＋ hashtag\n對齊 IG／FB 貼文需求")
    card(s, Inches(8.8), Inches(2.2), Inches(3.7), Inches(3.3), "私訊變現用", "懶人包整理\n留言領取的完整交付物", SOFT)
    footer(s, "結果句：產出物直接進工作流，不需要再人工拆欄。", 5)

    # 06
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), W, H, WHITE)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(10), Inches(0.35), "自定義規則", size=14, bold=True, color=GREEN_D)
    add_textbox(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1.1), "系統預設與客製化　留白沿用｜有填才改那一塊", size=28, bold=True, color=NAVY)
    card(
        s,
        Inches(0.7),
        Inches(2.2),
        Inches(5.8),
        Inches(3.4),
        "可調區塊",
        "身份｜語氣\n必要 hashtag 跟 CTA\n要／不要｜口播文案參考區\n禁止使用的題材\n自訂指令（進階，可啟用）",
        SOFT,
    )
    card(
        s,
        Inches(6.8),
        Inches(2.2),
        Inches(5.6),
        Inches(3.4),
        "增量覆寫",
        "只填語氣 → 只改口吻\nhashtag 留白 → 依題材長標籤\n空欄不會被亂補偏好\n\n底線仍保留：六欄／一次一角／繁中",
    )
    footer(s, "結果句：開箱通用，用久了變成你家帳號專屬產稿機。", 6)

    # 07 上手＋轉錄合併
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), W, H, WHITE)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(10), Inches(0.35), "上手一條龍", size=14, bold=True, color=GREEN_D)
    add_textbox(s, Inches(0.7), Inches(0.85), Inches(12), Inches(0.9), "打開網址就能產　有影片就先轉錄", size=30, bold=True, color=NAVY)
    card(s, Inches(0.55), Inches(2.2), Inches(2.9), Inches(3.3), "1｜打開產稿網址", "進入工作台\n不用安裝", SOFT)
    card(s, Inches(3.65), Inches(2.2), Inches(2.9), Inches(3.3), "2｜規則（可選）", "有要對齊再填\n沒有就全留白")
    card(s, Inches(6.75), Inches(2.2), Inches(2.9), Inches(3.3), "3｜影片丟轉錄（可選）", "網址進轉錄機器人\n複製文字", SOFT)
    card(s, Inches(9.85), Inches(2.2), Inches(2.9), Inches(3.3), "4｜貼上就產六欄", "貼稿／轉錄／題材\n選角度 → 貼 Sheet", NAVY, title_color=WHITE, body_color=RGBColor(0xC9, 0xE6, 0xD4))
    footer(s, "結果句：丟網址、貼文字、出稿——打開就能用。", 7)

    # 08
    s = prs.slides.add_slide(blank)
    add_rect(s, Inches(0), Inches(0), W, H, WHITE)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(10), Inches(0.35), "下一步", size=14, bold=True, color=GREEN_D)
    add_textbox(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1.1), "把產稿變成每天可交付的產線", size=30, bold=True, color=NAVY)
    card(s, Inches(0.55), Inches(2.3), Inches(2.9), Inches(3.1), "一打開就有", "雙模式產稿\n六欄輸出\n自定義規則\n禁題預警", SOFT)
    card(s, Inches(3.65), Inches(2.3), Inches(2.9), Inches(3.1), "含轉錄", "影片網址→文字\n再餵進產稿")
    card(s, Inches(6.75), Inches(2.3), Inches(2.9), Inches(3.1), "可擴", "多帳號／歷史\n規則雲端同步", SOFT)
    card(s, Inches(9.85), Inches(2.3), Inches(2.9), Inches(3.1), "核心不變", "結構保底\n題材判斷\n留言導向", NAVY, title_color=WHITE, body_color=RGBColor(0xC9, 0xE6, 0xD4))
    footer(s, "自動產稿機器人｜Short Script OS — 打開網址，就產。", 8)

    out = r"d:\各種好用的PY\自動產稿機器人\docs\自動產稿機器人_產品簡報.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
